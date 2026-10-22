"""Tests for the .docx loader used by the docs scanner.

The loader must work with the declared ``python-docx`` dependency only;
``docx2txt`` (what langchain's ``Docx2txtLoader`` needs) is not an AMX
dependency, so routing .docx through it silently dropped every Word
document at ingest time.
"""

from __future__ import annotations

import builtins
from pathlib import Path

import docx


def _make_docx(path: Path) -> Path:
    d = docx.Document()
    d.add_heading("Customer glossary", level=1)
    d.add_paragraph("KUND_NR is the customer number assigned at registration.")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "column"
    table.cell(0, 1).text = "meaning"
    table.cell(1, 0).text = "ORT"
    table.cell(1, 1).text = "city of the delivery address"
    d.save(path)
    return path


def test_docx_loader_returns_paragraphs_and_tables(tmp_path: Path) -> None:
    from amx.docs.loaders.docx_loader import load_docx

    text = load_docx(_make_docx(tmp_path / "glossary.docx"))
    assert "# Customer glossary" in text
    assert "KUND_NR is the customer number assigned at registration." in text
    assert "| column | meaning |" in text
    assert "| ORT | city of the delivery address |" in text


def test_docx_ingest_does_not_need_docx2txt(tmp_path: Path, monkeypatch) -> None:
    real_import = builtins.__import__

    def no_docx2txt(name, *args, **kwargs):
        if name == "docx2txt" or name.startswith("docx2txt."):
            raise ModuleNotFoundError("No module named 'docx2txt'")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", no_docx2txt)
    from amx.docs.rag import _build_loader_map

    loader = _build_loader_map()[".docx"](_make_docx(tmp_path / "sop.docx"))
    docs = loader.load()
    assert docs and "KUND_NR" in docs[0].page_content
